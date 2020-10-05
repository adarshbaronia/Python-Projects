
#reading data from a csv file which has dictionaries nested in list for the Keywords column
#writing that cleaned data to word doc

import pandas as pd
n=0
df=pd.read_csv('sample.csv')
s={}
for i, row in df.iterrows():
    if i==n:
        t=''
        if row['Keywords']!='':
            for j in row['Keywords'].strip('][').replace('{',"").replace('}',"").split(', '):
                if 'name' in j:
                    if row['ID'] in s.keys():
                        
                        s[row['ID']]=s[row['ID']]+","+j[9:].replace("'","")
                    
                    
                    else:
                        s[row['ID']]=","+j[9:].replace("'","")
                    
                    
        elif 'name' not in j:
            s[row['ID']]=''
    n+=1
#df=df.drop('names', axis=1)
ds=pd.DataFrame(s, index=range(len(s))).T
cols=ds.columns
ds = ds.drop(cols[1:], axis=1)
ds['ID']=ds.index
ds.columns=['Names', 'ID']
ds.Names=ds.Names.apply(lambda x:x[1:])

combined=df.merge(ds, on='ID', how='inner')

print(combined)

import docx

# open an existing document
doc = docx.Document('test.docx')

# add a table to the end and create a reference variable
# extra row is so we can add the header row
t = doc.add_table(combined.shape[0]+1, combined.shape[1])

# add the header rows.
for j in range(combined.shape[-1]):
    t.cell(0,j).text = combined.columns[j]

# add the rest of the data frame
for i in range(combined.shape[0]):
    for j in range(combined.shape[-1]):
        t.cell(i+1,j).text = str(combined.values[i,j])

# save the doc
doc.save('./test.docx')
print('Data has been written to word document!')

                

    
